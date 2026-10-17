# -*- coding: utf-8 -*-
"""
md2docx.py — Convierte un Markdown sencillo (el de docs/) a Word con estilo.

Soporta: # ## ### ####, párrafos con **negrita** y `código`, listas con - y 1.,
tablas GFM, bloques ``` ... ```, citas >, reglas --- e imágenes ![alt](ruta).
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = "1F6F8B"
ORANGE = "D2691E"
GREY = "595959"
LIGHT = "EAF3F6"
BODY_FONT = "Calibri"
MONO = "Consolas"


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def ancho_fijo(tabla, anchos):
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for fila in tabla.rows:
        for i, w in enumerate(anchos):
            if i < len(fila.cells):
                fila.cells[i].width = Cm(w)


def texto_con_formato(par, texto, base=10, color=None):
    """Añade el texto interpretando **negrita** y `código`."""
    partes = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto)
    for p in partes:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            r = par.add_run(p[2:-2])
            r.bold = True
        elif p.startswith("`") and p.endswith("`"):
            r = par.add_run(p[1:-1])
            r.font.name = MONO
            r.font.size = Pt(base - 0.5)
        else:
            r = par.add_run(p)
        r.font.size = Pt(base)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return par


def es_tabla(linea):
    return linea.strip().startswith("|") and linea.strip().endswith("|")


def celdas(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def construir(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lineas = f.read().split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)

    for nombre, tam, color in [("Heading 1", 17, TEAL), ("Heading 2", 13, TEAL),
                               ("Heading 3", 11, ORANGE)]:
        st = doc.styles[nombre]
        st.font.name = BODY_FONT
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True

    i = 0
    n = len(lineas)
    base_dir = os.path.dirname(os.path.abspath(md_path))
    titulo_puesto = False

    while i < n:
        linea = lineas[i]
        s = linea.strip()

        if not s:
            i += 1
            continue

        # --- imágenes
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            alt, ruta = m.group(1), m.group(2)
            ruta_abs = ruta if os.path.isabs(ruta) else os.path.join(base_dir, ruta)
            if os.path.exists(ruta_abs):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(ruta_abs, width=Cm(7.4))
                if alt:
                    c = doc.add_paragraph()
                    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = c.add_run(alt)
                    r.italic = True
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor.from_string(GREY)
            i += 1
            continue

        # --- reglas
        if re.match(r"^-{3,}$", s):
            i += 1
            continue

        # --- bloques de código
        if s.startswith("```"):
            i += 1
            bloque = []
            while i < n and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            i += 1
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            celda = t.rows[0].cells[0]
            sombrear(celda, LIGHT)
            celda.text = ""
            par = celda.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            r = par.add_run("\n".join(bloque))
            r.font.name = MONO
            r.font.size = Pt(8)
            ancho_fijo(t, [17.4])
            doc.add_paragraph()
            continue

        # --- tablas
        if es_tabla(s) and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lineas[i + 1].strip()):
            encabezados = celdas(s)
            i += 2
            filas = []
            while i < n and es_tabla(lineas[i].strip()):
                filas.append(celdas(lineas[i]))
                i += 1
            ncol = len(encabezados)
            tabla = doc.add_table(rows=1, cols=ncol)
            tabla.style = "Table Grid"
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(encabezados):
                celda = tabla.rows[0].cells[j]
                celda.text = ""
                par = celda.paragraphs[0]
                par.paragraph_format.space_after = Pt(1)
                r = par.add_run(h)
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor.from_string("FFFFFF")
                sombrear(celda, TEAL)
            for fila in filas:
                celdas_fila = tabla.add_row().cells
                for j in range(ncol):
                    celda = celdas_fila[j]
                    celda.text = ""
                    par = celda.paragraphs[0]
                    par.paragraph_format.space_after = Pt(1)
                    texto_con_formato(par, fila[j] if j < len(fila) else "", base=8.5)
            ancho_fijo(tabla, [17.4 / ncol] * ncol)
            doc.add_paragraph()
            continue

        # --- citas
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            texto_con_formato(p, s.lstrip("> ").strip(), base=9.5, color=GREY)
            r = p.runs[0] if p.runs else p.add_run("")
            r.italic = True
            i += 1
            continue

        # --- títulos
        if s.startswith("#"):
            nivel = len(s) - len(s.lstrip("#"))
            texto = s[nivel:].strip()
            if nivel == 1 and not titulo_puesto:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(texto)
                r.bold = True
                r.font.size = Pt(24)
                r.font.color.rgb = RGBColor.from_string(TEAL)
                titulo_puesto = True
            else:
                doc.add_heading(texto, level=min(max(nivel - 1, 1), 3))
            i += 1
            continue

        # --- listas
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            texto_con_formato(p, re.sub(r"^[-*]\s+", "", s), base=9.5)
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            texto_con_formato(p, m.group(2), base=9.5)
            i += 1
            continue

        # --- párrafo normal
        p = doc.add_paragraph()
        texto_con_formato(p, s, base=10)
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else "docs/Plan-App-LangKids-Android.md"
    salida = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(entrada)[0] + ".docx"
    construir(entrada, salida)
    print("OK ->", salida, os.path.getsize(salida) // 1024, "KB")
